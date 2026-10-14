"""Build the JCRAI 2026 title page in Word and PDF form.

JCRAI 2026 reviews anonymously: the manuscript goes to the reviewers without
author identity, and the author names, affiliations, contact details, and short
biographies travel on a separate title page.  The conference ships this title
page as a designed Word flyer whose content lives in text boxes; this script
emits an equivalent, plainly formatted document from one data table so the
English and Chinese versions cannot drift apart.

Outputs ``title_page_en.docx``/``title_page_zh.docx`` and the matching
``.tex`` sources, which compile with pdflatex and xetex respectively.
"""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt


ROOT = Path(__file__).resolve().parent

TITLE_EN = (
    "Typed Judgment, Deterministic Enforcement: A Consent-Gated Gateway "
    "Architecture for Inter-Agent Collaboration over Messaging Channels"
)
TITLE_ZH = "类型化判断与确定性执行：面向消息通道的同意门控智能体协作网关架构"

SESSION_EN = (
    "Special Session XII: Agent Interconnection and Collaboration Intelligence, "
    "JCRAI 2026"
)
SESSION_ZH = "JCRAI 2026 专题 XII：智能体互联与协作智能"

# Order follows the manuscript: Cai, Liu, Ma, Yu, Li.
AUTHORS = [
    {
        "en_name": "Yizhang Cai",
        "zh_name": "蔡亦张",
        "en_affil": (
            "School of Artificial Intelligence, Beijing University of Posts and "
            "Telecommunications, Beijing, China"
        ),
        "zh_affil": "北京邮电大学人工智能学院，中国北京",
        "email": "2025111259@bupt.cn",
        "en_bio": (
            "Yizhang Cai is a graduate student at the School of Artificial "
            "Intelligence, Beijing University of Posts and Telecommunications, "
            "Beijing, China. His research interests are message-channel agent "
            "gateways, typed decision models for routing and consent, and the "
            "security evaluation of agent interconnection protocols. He is the "
            "first author of this paper."
        ),
        "zh_bio": (
            "蔡亦张，北京邮电大学人工智能学院研究生，研究兴趣包括消息通道上的"
            "智能体网关、用于路由与同意的类型化决策模型，以及智能体互联协议的"
            "安全性评测。本文第一作者。"
        ),
    },
    {
        "en_name": "Jun Liu",
        "zh_name": "刘军",
        "en_affil": (
            "School of Artificial Intelligence, Beijing University of Posts and "
            "Telecommunications, Beijing, China"
        ),
        "zh_affil": "北京邮电大学人工智能学院，中国北京",
        "email": "liujun@bupt.edu.cn",
        "en_bio": (
            "Jun Liu is with the School of Artificial Intelligence, Beijing "
            "University of Posts and Telecommunications, Beijing, China. His "
            "research focuses on agent interconnection and collaboration "
            "protocols; he is the first author of the ACPs protocol suite for the "
            "Internet of Agents and a co-author of the ACPS protocol study. "
        ),
        "zh_bio": (
            "刘军，北京邮电大学人工智能学院教师，主要研究智能体互联与协作协议，"
            "是智能体互联网协议族 ACPs 的第一作者，也是 ACPS 协议研究的合作者。"
        ),
    },
    {
        "en_name": "Di Ma",
        "zh_name": "马镝",
        "en_affil": "Beijing Guoshu Zhilian Technology Co., Ltd., Beijing, China",
        "zh_affil": "北京国数智联科技有限公司，中国北京",
        "email": "madi@guoshuzhilian.com.cn",
        "en_bio": (
            "Di Ma is with Beijing Guoshu Zhilian Technology Co., Ltd., Beijing, "
            "China. He works on the engineering and deployment of AI systems and "
            "on the industrial requirements behind this study."
        ),
        "zh_bio": (
            "马镝，就职于北京国数智联科技有限公司，从事人工智能系统的工程化与"
            "部署工作，并为本文提供产业侧需求。"
        ),
    },
    {
        "en_name": "Ke Yu",
        "zh_name": "禹可",
        "en_affil": (
            "School of Artificial Intelligence, Beijing University of Posts and "
            "Telecommunications, Beijing, China"
        ),
        "zh_affil": "北京邮电大学人工智能学院，中国北京",
        "email": "yuke@bupt.edu.cn",
        "en_bio": (
            "Ke Yu is with the School of Artificial Intelligence, Beijing "
            "University of Posts and Telecommunications, Beijing, China. Her "
            "research covers network intelligence, traffic analysis, and agent "
            "collaboration protocols; she is a co-author of the ACPs protocol "
            "suite and of the ACPS protocol study."
        ),
        "zh_bio": (
            "禹可，北京邮电大学人工智能学院教师，研究领域包括网络智能、流量分析"
            "与智能体协作协议，是智能体互联网协议族 ACPs 与 ACPS 协议研究的"
            "合作者。"
        ),
    },
    {
        "en_name": "Ke Li",
        "zh_name": "李珂",
        "en_affil": (
            "School of Artificial Intelligence, Beijing University of Posts and "
            "Telecommunications, Beijing, China"
        ),
        "zh_affil": "北京邮电大学人工智能学院，中国北京",
        "email": "like1990@bupt.edu.cn",
        "en_bio": (
            "Ke Li is with the School of Artificial Intelligence, Beijing "
            "University of Posts and Telecommunications, Beijing, China, and is "
            "the corresponding author of this paper. His research interests are "
            "agent collaboration protocols, multi-agent systems, and protocol-level "
            "trust and security for the Internet of Agents."
        ),
        "zh_bio": (
            "李珂，北京邮电大学人工智能学院教师，本文通讯作者。研究兴趣包括"
            "智能体协作协议、多智能体系统，以及智能体互联网在协议层面的信任与"
            "安全。"
        ),
    },
]


def build_docx(lang: str) -> Path:
    """Write the Word title page for ``lang`` in {"en", "zh"}."""
    zh = lang == "zh"
    title = TITLE_ZH if zh else TITLE_EN
    session = SESSION_ZH if zh else SESSION_EN
    headers = ("作者", "简介") if zh else ("Name", "Brief Bio")
    out = ROOT / f"title_page_{lang}.docx"

    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(11)

    heading = doc.add_paragraph()
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = heading.add_run("标题页（Title Page）" if zh else "Title Page")
    run.bold = True
    run.font.size = Pt(14)

    para = doc.add_paragraph()
    run = para.add_run(title)
    run.bold = True

    para = doc.add_paragraph()
    para.add_run(("专题：" if zh else "Session: ") + session)

    doc.add_paragraph()
    head = doc.add_paragraph()
    head.add_run("作者信息" if zh else "Author details").bold = True
    for author in AUTHORS:
        para = doc.add_paragraph(style="List Bullet")
        name = author["zh_name"] if zh else author["en_name"]
        affil = author["zh_affil"] if zh else author["en_affil"]
        para.add_run(f"{name} — {affil} — {author['email']}")

    doc.add_paragraph()
    head = doc.add_paragraph()
    head.add_run("作者简介" if zh else "Author biographies").bold = True
    table = doc.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    for cell, text in zip(table.rows[0].cells, headers):
        cell.text = ""
        cell.paragraphs[0].add_run(text).bold = True
    for author in AUTHORS:
        row = table.add_row().cells
        name = author["zh_name"] if zh else author["en_name"]
        bio = author["zh_bio"] if zh else author["en_bio"]
        row[0].text = name
        row[1].text = bio

    doc.save(out)
    print(out)
    return out


def build_tex(lang: str) -> Path:
    """Write a LaTeX rendering of the same title page."""
    zh = lang == "zh"
    title = TITLE_ZH if zh else TITLE_EN
    session = SESSION_ZH if zh else SESSION_EN
    out = ROOT / f"title_page_{lang}.tex"

    lines = []
    if zh:
        lines += [
            r"\documentclass[11pt,a4paper]{article}",
            r"\usepackage{geometry}",
            r"\geometry{left=2.2cm,right=2.2cm,top=2.2cm,bottom=2.2cm}",
            r"\usepackage{fontspec}",
            r"\setmainfont{Noto Serif CJK SC}",
            r'\XeTeXlinebreaklocale "zh"',
            r"\XeTeXlinebreakskip = 0pt plus 1pt minus 0.1pt",
            r"\setlength{\parindent}{0pt}",
            r"\setlength{\parskip}{0.6ex}",
        ]
    else:
        lines += [
            r"\documentclass[11pt,a4paper]{article}",
            r"\usepackage[T1]{fontenc}",
            r"\usepackage{newtxtext}",
            r"\usepackage{geometry}",
            r"\geometry{left=2.2cm,right=2.2cm,top=2.2cm,bottom=2.2cm}",
            r"\setlength{\parindent}{0pt}",
            r"\setlength{\parskip}{0.6ex}",
        ]
    lines += [
        r"\usepackage{longtable}",
        r"\usepackage{array}",
        r"\begin{document}",
        r"\begin{center}{\Large\bfseries " + tex_escape(title) + r"}\\[0.6em]",
        tex_escape(session),
        r"\end{center}",
        r"\vspace{1em}",
        r"{\bfseries " + ("作者信息" if zh else "Author details") + r"}",
        r"\begin{itemize}\setlength{\itemsep}{0.1em}\setlength{\parskip}{0em}",
    ]
    for author in AUTHORS:
        name = author["zh_name"] if zh else author["en_name"]
        affil = author["zh_affil"] if zh else author["en_affil"]
        lines.append(
            r"\item " + tex_escape(name) + ", " + tex_escape(affil)
            + r", \texttt{" + tex_escape(author["email"]) + r"}"
        )
    lines += [
        r"\end{itemize}",
        r"\vspace{1em}",
        r"{\bfseries " + ("作者简介" if zh else "Author biographies") + r"}",
        r"\begin{longtable}{@{}p{3.4cm}p{10.8cm}@{}}",
        r"\hline",
        r"\textbf{" + ("作者" if zh else "Name") + r"} & \textbf{"
        + ("简介" if zh else "Brief Bio") + r"} \\[0.3em]",
        r"\hline",
        r"\endhead",
    ]
    for author in AUTHORS:
        name = author["zh_name"] if zh else author["en_name"]
        bio = author["zh_bio"] if zh else author["en_bio"]
        lines.append(
            r"{\bfseries " + tex_escape(name) + r"} & {\footnotesize " + tex_escape(bio)
            + r"} \\[0.4em] \hline"
        )
    lines += [r"\end{longtable}", r"\end{document}"]
    out.write_text("\n".join(lines) + "\n", encoding="utf-8")
    print(out)
    return out


def tex_escape(text: str) -> str:
    for char, repl in (
        ("\\", r"\textbackslash{}"),
        ("&", r"\&"),
        ("%", r"\%"),
        ("$", r"\$"),
        ("#", r"\#"),
        ("_", r"\_"),
        ("{", r"\{"),
        ("}", r"\}"),
        ("~", r"\textasciitilde{}"),
        ("^", r"\textasciicircum{}"),
    ):
        text = text.replace(char, repl)
    return text


if __name__ == "__main__":
    for lang in ("en", "zh"):
        build_docx(lang)
        build_tex(lang)
